from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(16, 42, 67)
BLUE = RGBColor(71, 96, 255)
BLUE_DARK = RGBColor(47, 66, 190)
TEAL = RGBColor(0, 151, 167)
GREEN = RGBColor(23, 143, 94)
ORANGE = RGBColor(231, 122, 58)
GRAY = RGBColor(91, 107, 121)
LIGHT_BLUE = RGBColor(237, 241, 255)
LIGHT_TEAL = RGBColor(232, 248, 248)
LIGHT_GREEN = RGBColor(234, 247, 240)
LIGHT_ORANGE = RGBColor(255, 244, 234)
LIGHT_GRAY = RGBColor(244, 246, 248)
WHITE = RGBColor(255, 255, 255)


def new_deck() -> Presentation:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    return deck


def add_blank_slide(deck: Presentation):
    return deck.slides.add_slide(deck.slide_layouts[6])


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def text_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float = 14,
    color: RGBColor = NAVY,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin: float = 0.08,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    heading: str,
    detail: str = "",
    fill: RGBColor = LIGHT_BLUE,
    line: RGBColor = BLUE,
    heading_color: RGBColor = NAVY,
    detail_color: RGBColor = GRAY,
    heading_size: float = 15,
    detail_size: float = 10.5,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.4)
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.12)
    frame.margin_top = frame.margin_bottom = Inches(0.07)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = heading
    run.font.name = "Aptos"
    run.font.size = Pt(heading_size)
    run.font.bold = True
    run.font.color.rgb = heading_color
    if detail:
        paragraph = frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_before = Pt(2)
        run = paragraph.add_run()
        run.text = detail
        run.font.name = "Aptos"
        run.font.size = Pt(detail_size)
        run.font.color.rgb = detail_color
    return shape


def arrow(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    color: RGBColor = GRAY,
    width: float = 1.5,
    dashed: bool = False,
):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    connector.line.end_arrowhead = True
    if dashed:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return connector


def rule(
    slide,
    x: float,
    y: float,
    w: float,
    color: RGBColor = BLUE,
    height: float = 0.035,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def title(slide, heading: str, kicker: str, subtitle: str) -> None:
    text_box(slide, 0.68, 0.26, 12.0, 0.28, kicker.upper(), 9, BLUE, True)
    text_box(slide, 0.68, 0.52, 12.0, 0.58, heading, 27, NAVY, True)
    text_box(slide, 0.70, 1.05, 11.9, 0.34, subtitle, 11.5, GRAY)


def link(slide, text: str, url: str):
    shape = text_box(
        slide,
        9.35,
        7.05,
        3.25,
        0.22,
        text,
        8.5,
        TEAL,
        align=PP_ALIGN.RIGHT,
    )
    shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    return shape


def place_picture(
    slide,
    source: str | Path | bytes,
    x: float,
    y: float,
    max_w: float,
    max_h: float,
):
    image = BytesIO(source) if isinstance(source, bytes) else str(source)
    shape = slide.shapes.add_picture(
        image,
        Inches(x),
        Inches(y),
        width=Inches(max_w),
    )
    if shape.height > Inches(max_h):
        ratio = shape.height / shape.width
        shape.height = Inches(max_h)
        shape.width = Emu(int(Inches(max_h) / ratio))
    shape.left = Inches(x) + Emu(int((Inches(max_w) - shape.width) / 2))
    shape.top = Inches(y) + Emu(int((Inches(max_h) - shape.height) / 2))
    return shape
